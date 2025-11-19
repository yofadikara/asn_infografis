from dataconfig.data import get_all_data, get_table_name, tambah_persentase, klasifikasi_kelompok_jabatan, ringkasan_kelompok_jabatan
from dataconfig.db_config import connect_db, load_env
from automation.export_excel import export_per_provinsi, export_nasional
from automation.export_ppt import isi_infografis, generate_instansi_values, isi_balok_pendidikan, generate_all_tables, generate_pie_chart
from pptx import Presentation
import pandas as pd
from dataconfig.mapping import mapping_placeholder_per_kategori, label_column_per_kategori, selid_kategori_shared
from io import BytesIO
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseUpload
from template.query_templates import get_instansi_list
import os, sys
from automation.filter import get_instansi_filter, resolve_filter

#TARIK DATA
params = load_env()
conn = connect_db(params)
args = get_instansi_filter()
#MODE NASIONAL
if args.nasional:
    mode = 'nasional'
    target_list = [{}]
    nama_table = get_table_name()
    print("\nNama tabel yang digunakan:", nama_table,"\n")
    data_nasional = get_all_data(conn, nama_table, target_list, mode=mode)
    prs = Presentation('template/template_infografis_nasional.pptx')
    print("Memproses Data Nasional")
    kategori_items = list(data_nasional["Nasional"].items())
    for kategori, df in kategori_items:
        if 'count' in df.columns:
            #Penambahan Kolom Total pada Jenis ASN
            if 'jenis_asn' in df.columns and 'TOTAL' not in df['jenis_asn'].values:
                df_total = pd.DataFrame({
                    'jenis_asn': ['TOTAL'],
                    'count': [df['count'].sum()]
                })
                df = pd.concat([df, df_total], ignore_index=True)
                df_clean = df[df['jenis_asn'] != 'TOTAL'].copy()
                df_clean = tambah_persentase(df_clean) #tambah persentase
                df_final = pd.concat([
                    df_clean,
                    df[df['jenis_asn'] == 'TOTAL']
                ], ignore_index=True)
                data_nasional["Nasional"][kategori] = df_final
            #Penambahan Ringkasan Kelompok Jabatan       
            elif 'jenisjabatan' in df.columns and 'count' in df.columns:
                df = klasifikasi_kelompok_jabatan(df)
                #data_dict[prov][kategori] = tambah_persentase(df) 
                df_ringkasan = ringkasan_kelompok_jabatan(df)
                data_nasional["Nasional"][f"Ringkasan_{kategori}"] = df_ringkasan
            else:
                data_nasional["Nasional"][kategori] = tambah_persentase(df)
    #TABLE DINAMIS
    generate_all_tables(prs, data_nasional["Nasional"], mode)
    kategori_items = list(data_nasional["Nasional"].items())
    for kategori, df in kategori_items:
        if kategori == "Pendidikan":
            slide_index = selid_kategori_shared.get(kategori)
            slide = prs.slides[slide_index]
            df_export = data_nasional["Nasional"][kategori]
            isi_balok_pendidikan(df_export, slide, mode)
        if kategori == "Kelompok_Generasi":
            slide_index = selid_kategori_shared.get(kategori)
            slide = prs.slides[slide_index]
            df_export = data_nasional["Nasional"][kategori]
            generate_pie_chart(
                df_export,
                label_column="kelompok_generasi",
                value_column="count",
                slide=slide,
                left_cm=18.37,
                top_cm=3.38,
                width_cm=6.65,
                height_cm=4.24
            )
        #Isi Data PPT
        if kategori in mapping_placeholder_per_kategori:
            mapping_placeholder = mapping_placeholder_per_kategori[kategori]
            label_column = label_column_per_kategori.get(kategori, "label")
            #Nama Instansi Dinamis                  
            mapping_instansi = generate_instansi_values('Nasional')
            data_export = data_nasional["Nasional"][kategori]
            if data_export is not None:
                print(f"Memproses isi_infografis untuk kategori: {kategori}")
                isi_infografis(prs, data_export, mapping_placeholder, mapping_instansi, label_column= label_column)
            else:
                print(f"Data Kosong untuk Nasional, Kategori: {kategori}")
    #Save PPT
    output_path = "output/ppt/Statistik_Nasional.pptx"
    prs.save(output_path)
    print(f"PowerPoint Disimpan: {output_path}")
    #Export ke Excel
    if data_nasional and any(data_nasional.values()):
        export_per_provinsi(data_nasional)
    else:
        print("Tidak ada data yang tersedia untuk diekspor.")
    sys.exit()
    
#MODE INSTANSI
#Mengambil Prefix Wilker
filter_value = resolve_filter(args, conn)
#Kondisi Mode
if args.instansi:
    mode = 'instansi'
    target_list = get_instansi_list(conn, filter_value)
elif args.provinsi:
    mode = 'provinsi'
    target_list = filter_value
elif args.kanreg:
    mode = 'kanreg'
    target_list = filter_value
else:
    raise ValueError("Mode harus 'instansi', 'provinsi', atau 'kanreg'")
nama_table = get_table_name()
print("\nNama tabel yang digunakan:", nama_table,"\n")
data_dict = get_all_data(conn, nama_table, target_list or [], mode=mode)
for label, kategori_data in data_dict.items():
    prs = Presentation('template/template_infografis.pptx')
    print(f"Data untuk Provinsi: {label}")
    kategori_items = list(kategori_data.items()) #untuk mengindari runtime error saat iterasi
    for kategori, df in kategori_items:
        if 'count' in df.columns:
            #Penambahan Kolom Total pada Jenis ASN
            if 'jenis_asn' in df.columns and 'TOTAL' not in df['jenis_asn'].values:
                df_total = pd.DataFrame({
                    'jenis_asn': ['TOTAL'],
                    'count': [df['count'].sum()]
                })
                df = pd.concat([df, df_total], ignore_index=True)
                df_clean = df[df['jenis_asn'] != 'TOTAL'].copy()
                df_clean = tambah_persentase(df_clean) #tambah persentase
                df_final = pd.concat([
                    df_clean,
                    df[df['jenis_asn'] == 'TOTAL']
                ], ignore_index=True)
                data_dict[label][kategori] = df_final
            #Penambahan Ringkasan Kelompok Jabatan       
            elif 'jenisjabatan' in df.columns and 'count' in df.columns:
                df = klasifikasi_kelompok_jabatan(df)
                #data_dict[prov][kategori] = tambah_persentase(df) 
                df_ringkasan = ringkasan_kelompok_jabatan(df)
                data_dict[label][f"Ringkasan_{kategori}"] = df_ringkasan
            else:
                data_dict[label][kategori] = tambah_persentase(df)
    #TABLE DINAMIS
    generate_all_tables(prs, data_dict[label])
    #Add Ulang Kategori           
    kategori_items = list(kategori_data.items())
    for kategori, df in kategori_items:
        if kategori == "Pendidikan":
            slide_index = selid_kategori_shared.get(kategori)
            slide = prs.slides[slide_index]
            df_export = data_dict[label][kategori]
            isi_balok_pendidikan(df_export, slide, mode)
        #Isi Data PPT
        if kategori in mapping_placeholder_per_kategori:
            mapping_placeholder = mapping_placeholder_per_kategori[kategori]
            label_column = label_column_per_kategori.get(kategori, "label")
            #Nama Instansi Dinamis                  
            mapping_instansi = generate_instansi_values(label)
            data_export = data_dict[label][kategori]
            if data_export is not None:
                '''print(f"Memproses isi_infografis untuk kategori: {kategori}")'''
                isi_infografis(prs, data_export, mapping_placeholder, mapping_instansi, label_column= label_column)
            else:
                print(f"Data Kosong untuk Provinsi: {label}, Kategori: {kategori}")
    #Save PPT
    output_path = f"output/ppt/{label}.pptx"
    prs.save(output_path)
    print(f"PowerPoint Disimpan: {output_path}")
    #Upload Drive
    '''SCOPES = ['https://www.googleapis.com/auth/drive.file']
    TOKEN_FILE = 'token.json'
    CREDENTIALS_FILE = 'credentials.json'
    if os.path.exists(TOKEN_FILE):
        creds = Credentials.from_authorized_user_file(TOKEN_FILE, SCOPES)
    else:
        flow = InstalledAppFlow.from_client_secrets_file(CREDENTIALS_FILE, SCOPES)
        creds = flow.run_local_server(port=0)
        with open(TOKEN_FILE, 'w') as token:
            token.write(creds.to_json())
    drive_service = build('drive', 'v3', credentials=creds)
    #save to bytes sementara
    ppt_stream = BytesIO()
    prs.save(ppt_stream)
    ppt_stream.seek(0)
    #Upload
    file_metadata = {
        'name': f'{prov}.pptx',
        'parents': ['1OhAKH94lYQwZVAarrz1nlWW86X3aBKJv']
    }
    media = MediaIoBaseUpload(ppt_stream, mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation',
                              resumable=True)
    file = drive_service.files().create(body=file_metadata, media_body=media, fields='id').execute()
    print(f"File {prov}.pptx diunggah ke Google Drive dengan ID: {file.get('id')}")
'''
    '''#Export ke Excel
    if data_dict and any(data_dict.values()):
        export_per_provinsi(data_dict)
    else:
        print("Tidak ada data yang tersedia untuk diekspor.")
        '''