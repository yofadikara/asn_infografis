from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE 
from pptx.enum.text import PP_ALIGN
import matplotlib.pyplot as plt
import pandas as pd
import re
from io import BytesIO
from dataconfig.data import tambah_persentase, get_tanggal_data
from dataconfig.mapping import mapping_diagram_cm, mapping_textbox_cm, mapping_diagram_nasional, mapping_textbox_nasional, mapping_color_textbox

#fungsi untuk mengganti teks pada shape dengan mempertahankan gaya
def replace_text_reserving_style(shape, placeholder, replacement):
    if not shape.has_text_frame:
        return
    for paragraph in shape.text_frame.paragraphs:
        full_text = ''.join([run.text for run in paragraph.runs])
        if placeholder not in full_text:
                continue
        replaced_text = full_text.replace(placeholder, replacement)
        for run in paragraph.runs:
                run.text = ''
                if paragraph.runs:
                     paragraph.runs[0].text = replaced_text

def generate_instansi_values(prov):
     return{
          "{instansi_1}": prov,
          "{tanggal_data}": get_tanggal_data()
     }

#fungsi isi placeholder infografis
def isi_infografis(prs, df_kategori, mapping_placeholder, mapping_instansi, label_column='label'):
    slide = prs.slides[0]
    #untuk penjagaan placeholder kosong
    used_placeholders = set()
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for placeholder, replacement in mapping_instansi.items():
             replace_text_reserving_style(shape, placeholder, replacement)
             used_placeholders.add(placeholder)
        text = shape.text
        for idx, row in df_kategori.iterrows():
            if label_column == "jabatan":
                kelompok = row.get("kelompok_jabatan")
                jabatan = row.get("jabatan")
                placeholder = mapping_placeholder["count"].get(kelompok,{}).get(jabatan,None)
                if placeholder:
                    nilai = row.get("count")
                    nilai_str = f"{nilai:,}".replace(",",".") if isinstance(nilai, (int)) else str(nilai)
                    replace_text_reserving_style(shape, placeholder, nilai_str)
                    used_placeholders.add(placeholder)
            else:
                label = row.get(label_column)
                for kolom_data, map_label in mapping_placeholder.items():
                    if label in map_label:
                        placeholder = map_label[label]
                        nilai = row.get(kolom_data)
                        nilai_str = f"{nilai:,}".replace(",",".") if isinstance(nilai, (int)) else str(nilai)
                        replace_text_reserving_style(shape, placeholder, nilai_str)
                        used_placeholders.add(placeholder)
    #penjagaan placeholder yang tidak terpakai
    all_placeholders = set(mapping_instansi.keys())
    if label_column == "jabatan":
        for kelompok_dict in mapping_placeholder.get("count",{}).values():
            for v in kelompok_dict.values():
                if isinstance(v, str):
                    all_placeholders.add(v)
    else:
        for map_label in mapping_placeholder.values():
            all_placeholders.update(map_label.values())
    unused_placeholders = all_placeholders - used_placeholders
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for placeholder in unused_placeholders:
            if "persentase" in placeholder.lower() or "percent" in placeholder.lower():
                default_value = "0%"
            else:
                default_value = "0"
            replace_text_reserving_style(shape, placeholder, default_value)

#konversi cm ke inches
def cm(val):
    return Inches(val / 2.54)

#Fungsi Balok Pendidikan
def isi_balok_pendidikan(df,slide,mode=""):
    max_val = df['persentase_vis_num'].max()
    max_height = cm(5)
    mapping_diagram = mapping_diagram_nasional if mode == "nasional" else mapping_diagram_cm
    mapping_text = mapping_textbox_nasional if mode == "nasional" else mapping_textbox_cm
    for _, row in df.iterrows():
         label = row['tingkatpendidikan']
         count = row['count']
         persen_vis = row['persentase_vis_num']
         if label in mapping_diagram:
            pos = mapping_diagram[label]
            base_y = cm(pos['y'])
            tinggi = (persen_vis / max_val) * max_height
            if tinggi > base_y:
                 tinggi = base_y - cm(0.1)  # mengurangi sedikit agar tidak melebihi batas
            y_top = base_y - tinggi
            #Balok Diagram
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                cm(pos['x']), y_top,
                cm(pos['width']), tinggi
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = pos['color']
            shape.line.fill.background()  # hilangkan border
            #Textbox Nilai di Atas Balok
            base_y = cm(pos['y'])
            text_pos = mapping_text[label]
            textbox = slide.shapes.add_textbox(
                cm(text_pos['x']),
                base_y - tinggi - cm(text_pos['y_offset']),
                cm(text_pos['width']),
                cm(1.0)
            )
            tf = textbox.text_frame
            tf.text = f"{persen_vis:.0f}%\n{count:,}".replace(",",".")
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            tf.paragraphs[1].alignment = PP_ALIGN.CENTER
            #baris persentase
            run0 = tf.paragraphs[0].runs[0]
            font0 = run0.font
            font0.size = Pt(18) if mode == "nasional" else Pt(20)
            font0.bold = True
            font0.name = "Open Sans Bold"
            if label in mapping_color_textbox:
                font0.color.rgb = mapping_color_textbox[label]
            else:
                font0.color.rgb = RGBColor(89,89,89)
            #baris count
            run1 = tf.paragraphs[1].runs[0]
            font1 = run1.font
            font1.size = Pt(11.5) if mode == "nasional" else Pt(14)
            font1.name = "Open Sans Bold"
            font1.color.rgb = RGBColor(89,89,89)
            
def hex_to_rgb(hex_color):
    hex_color = hex_color.lstrip('#')
    return RGBColor(int(hex_color[0:2], 16), 
                    int(hex_color[2:4], 16), 
                    int(hex_color[4:6], 16))    

def sanitize_filename(name: str) -> str:
    # Ganti karakter ilegal dengan underscore atau spasi
    return re.sub(r'[<>:"/\\|?*]', '_', name)
        
#pembuatan table dinamis
'''def generate_table_from_data(slide, df, label_column, value_column, pos_x_cm,
                             pos_y_cm, col_widths_cm, margin, fill_color=None,
                             font_size_override=None):
     df_filtered = df[df[value_column] > 0].copy()
     if df_filtered.empty:
          print("Data kosong, tabel tidak dibuat.")
          return
     row_count = len(df_filtered)
     col_count = 2
     row_height = 0.27 #tinggi per baris
     table_height = row_height * row_count
     table_shape = slide.shapes.add_table(
         row_count, col_count,
         cm(pos_x_cm), cm(pos_y_cm),
         cm(sum(col_widths_cm)), cm(table_height)
     )
     table = table_shape.table
     #set lebar kolom
     for idx, width in enumerate(col_widths_cm):
         table.columns[idx].width = cm(width)
     table.first_row = False
     #set margin
     for row in table.rows:
         for cell in row.cells:
             cell.margin_top = cm(margin.get("top", 0.1))
             cell.margin_bottom = cm(margin.get("bottom", 0.1))
             cell.margin_left = cm(margin.get("left", 0.1))
             cell.margin_right = cm(margin.get("right", 0.1))
     #isi data
     for i, (_, row) in enumerate(df_filtered.iterrows()):
        #label_text = row[label_column]
        label_text = " ".join(str(cell) for cell in row)
        is_pppk = "PPPK" in label_text
        is_fungsional = "Fungsional" in label_text
        #pewarnaan sesuai mapping
        for j in range(col_count):
            cell = table.cell(i, j)
            #pewarnaan PPK
            if is_pppk and is_fungsional:
                cell.fill.solid()
                cell.fill.fore_color.rgb = hex_to_rgb("D0E0E3")
            elif fill_color:
                cell.fill.solid()
                cell.fill.fore_color.rgb = hex_to_rgb(fill_color)
        #kategori
        table.cell(i, 0).text = row[label_column]
        table.cell(i,0).text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        for run in table.cell(i, 0).text_frame.paragraphs[0].runs:
            font = run.font
            font.size = Pt(font_size_override) if font_size_override else Pt(9)
        #value
        value = int(row[value_column]) if pd.notnull(row[value_column]) else 0
        table.cell(i, 1).text = f"{value:,}".replace(",",".")
        table.cell(i, 1).text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
        for run in table.cell(i, 1).text_frame.paragraphs[0].runs:
            font = run.font
            font.size = Pt(font_size_override) if font_size_override else Pt(9)'''

#fungsi sort table
'''def sort_table_by_mapping(df, label_column, value_column, label_order):
    available_labels = df[label_column].unique().tolist()
    filtered_order = [label for label in label_order if label in available_labels]
    df_sorted = (
        df[df[label_column].isin(filtered_order)]
        .set_index(label_column)
        .reindex(filtered_order)
        .dropna(subset=[value_column])
        .reset_index()
    )
    return df_sorted'''

#fungsi generate all table dinamis
'''def generate_all_tables(prs, data_all, mode=""):
    config_table = TABLE_POSITIONS_NAS if mode == "nasional" else TABLE_POSITIONS
    for kategori, config in config_table.items():
        if kategori not in data_all:
            continue
        if kategori == "Jenis_Jabatan":
            df_jabatan = data_all.get(kategori)
            if df_jabatan is None or "kelompok_jabatan" not in df_jabatan.columns:
                continue
            for kelompok, sub_config in config.items():
                df_sub_raw = df_jabatan[df_jabatan["kelompok_jabatan"] == kelompok]
                #sort sesuai mapping
                label_order = sub_config.get("label_order")
                df_sub = sort_table_by_mapping(
                    df_sub_raw,
                    label_column=sub_config["label_column"],
                    value_column=sub_config["value_column"],
                    label_order=label_order
                ) if label_order else df_sub_raw
                if df_sub.empty:
                    continue
                slide = prs.slides[sub_config["slide_index"]]
                font_size = 7 if kelompok == "Struktural" and mode == "nasional" else 6
                generate_table_from_data(
                    slide=slide,
                    df=df_sub,
                    label_column=sub_config["label_column"],
                    value_column=sub_config["value_column"],
                    pos_x_cm=sub_config["pos_x_cm"],
                    pos_y_cm=sub_config["pos_y_cm"],
                    col_widths_cm=sub_config["col_widths_cm"],
                    margin=sub_config["margin"],
                    fill_color=sub_config.get("fill_color"),
                    font_size_override=font_size
                )
        else:
            df_kategori = data_all[kategori]
            font_size = None
            generate_table_from_data(
                slide=prs.slides[config["slide_index"]],
                df=df_kategori,
                label_column=config["label_column"],
                value_column=config["value_column"],
                pos_x_cm=config["pos_x_cm"],
                pos_y_cm=config["pos_y_cm"],
                col_widths_cm=config["col_widths_cm"], 
                margin=config["margin"],
                fill_color=config.get("fill_color"),
                font_size_override=font_size
            )
    if mode != "nasional":
        for kategori, config in VIS_TABLE_POSITIONS.items():
            if kategori not in data_all:
                continue
            df_vis = data_all[kategori]
            slide = prs.slides[config["slide_index"]]
            value_column = config["value_column"]
            pos_x_cm = config["pos_x_cm"]
            pos_y_cm = config["pos_y_cm"]
            col_widths_cm = config["col_widths_cm"]
            margin = config["margin"]
            fill_color = config.get("fill_color")
            generate_vis_table(
                slide=slide,
                df=df_vis,
                value_column=value_column,
                pox_x_cm=pos_x_cm,
                pos_y_cm=pos_y_cm,
                col_widths_cm=col_widths_cm[0],
                margin=margin,
                fill_color=fill_color
            )'''

#fungsi table transparan persentase
'''def generate_vis_table(slide, df, value_column, pox_x_cm,
                       pos_y_cm, col_widths_cm, margin, fill_color=None):
    df_filtered = df[df[value_column].str.rstrip('%').astype(float) > 0].copy()
    if df_filtered.empty:
        return
    row_count = len(df_filtered)
    table_shape = slide.shapes.add_table(
        row_count, 1,
        cm(pox_x_cm), cm(pos_y_cm),
        cm(col_widths_cm), cm(0.8)
    )
    table = table_shape.table
    #set margin
    for row in table.rows:
        for cell in row.cells:
            cell.margin_top = cm(margin.get("top", 0.1))
            cell.margin_bottom = cm(margin.get("bottom", 0.1))
            cell.margin_left = cm(margin.get("left", 0.1))
            cell.margin_right = cm(margin.get("right", 0.1))
    for i, (_, row) in enumerate(df_filtered.iterrows()):
        cell = table.cell(i, 0)
        cell.text = row[value_column]
        paragraph = cell.text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.RIGHT
        for run in paragraph.runs:
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor(31,78,121)
        #set fill color
        if fill_color:
            cell.fill.solid()
            cell.fill.fore_color.rgb = hex_to_rgb(fill_color)
        else:
            cell.fill.background()'''

#fungsi generate piechart
'''def generate_pie_chart(df, label_column, value_column, slide, left_cm, top_cm, width_cm=6, height_cm=6):
    #Siapkan data
    labels = [f"{label}\n({persen}%)" for label, persen in zip(df[label_column], df["persentase_vis"])]
    sizes = df[value_column]
    colors = [color_map[label] for label in df[label_column]]
    #Buat pie chart
    fig, ax = plt.subplots()
    wedges, texts = ax.pie(
        sizes,
        labels=labels,
        startangle=90,
        textprops={'fontsize': 11},
        colors=colors
    )
    ax.axis('equal')  # Pie chart bulat
    #Simpan ke buffer
    img_stream = BytesIO()
    plt.savefig(img_stream, format='png', bbox_inches='tight', transparent=True)
    plt.close(fig)
    img_stream.seek(0)
    # Masukkan ke slide
    slide.shapes.add_picture(
        img_stream,
        cm(left_cm),
        cm(top_cm),
        width=cm(width_cm),
        height=cm(height_cm)
    )'''

