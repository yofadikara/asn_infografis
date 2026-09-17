from dataconfig.data import get_all_data, get_table_name, tambah_persentase, ringkasan_kelompok_jabatan
from dataconfig.db_config import connect_db, load_env
from automation.export_excel import export_per_provinsi, export_nasional
from automation.export_ppt import isi_infografis, generate_instansi_values, isi_balok_pendidikan, sanitize_filename
from pptx import Presentation
import pandas as pd
from dataconfig.mapping import mapping_placeholder_per_kategori, label_column_per_kategori, selid_kategori_shared, logger
from io import BytesIO
'''from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseUpload'''
from template.query_templates import get_instansi_list, get_all_instansi, get_provinsi_prefix, get_kanreg_prefix, get_all_provinsi
import os, sys
import subprocess
from automation.filter import resolve_filter, normalize_args
from datetime import datetime
from automation.s3_integration import upload_file_s3, save_to_db

#TARIK DATA
def run_pipeline(args):
    args = normalize_args(args)
    params = load_env()
    conn = connect_db(params)
    '''args = get_instansi_filter()'''
    #MODE NASIONAL
    if args.get("nasional"):
        mode = 'nasional'
        target_list = [{}]
        nama_table = get_table_name()
        logger.info(f"Nama tabel yang digunakan: {nama_table}")
        data_nasional = get_all_data(conn, nama_table, target_list, mode=mode)
        prs = Presentation('template/template_infografis_nasional.pptx')
        logger.info("Memproses Data Nasional")
        kategori_items = list(data_nasional["Nasional"].items())
        for kategori, df in kategori_items:
            if 'count' in df.columns:
            #Penambahan Kolom Total pada Jenis ASN
                if 'jenis_asn' in df.columns and 'TOTAL' not in df['jenis_asn'].values:
                    df_total = pd.DataFrame({
                    'jenis_asn': ['TOTAL'],
                    'count': [df['count'].sum()]
                })
                    df = pd.concat([df, df_total], ignore_index=True)
                    df_clean = df[df['jenis_asn'] != 'TOTAL'].copy()
                    df_clean = tambah_persentase(df_clean) #tambah persentase
                    df_final = pd.concat([
                        df_clean,
                        df[df['jenis_asn'] == 'TOTAL']
                    ], ignore_index=True)
                    data_nasional["Nasional"][kategori] = df_final
            #Penambahan Ringkasan Kelompok Jabatan       
                elif 'kelompok_jabatan' in df.columns and 'count' in df.columns:
                    '''df = klasifikasi_kelompok_jabatan(df)
                    data_dict[prov][kategori] = tambah_persentase(df)''' 
                    df_ringkasan = ringkasan_kelompok_jabatan(df)
                    data_nasional["Nasional"][f"Ringkasan_{kategori}"] = df_ringkasan
                else:
                    data_nasional["Nasional"][kategori] = tambah_persentase(df)
        #TABLE DINAMIS
        '''generate_all_tables(prs, data_nasional["Nasional"], mode)'''
        kategori_items = list(data_nasional["Nasional"].items())
        for kategori, df in kategori_items:
            if kategori == "Pendidikan":
                slide_index = selid_kategori_shared.get(kategori)
                slide = prs.slides[slide_index]
                df_export = data_nasional["Nasional"][kategori]
                isi_balok_pendidikan(df_export, slide, mode)
            #Isi Data PPT
            if kategori in mapping_placeholder_per_kategori:
                mapping_placeholder = mapping_placeholder_per_kategori[kategori]
                label_column = label_column_per_kategori.get(kategori, "label")
                #Nama Instansi Dinamis                  
                mapping_instansi = generate_instansi_values('Nasional')
                data_export = data_nasional["Nasional"][kategori]
                if data_export is not None:
                    '''print(f"Memproses isi_infografis untuk kategori: {kategori}")'''
                    isi_infografis(prs, data_export, mapping_placeholder, mapping_instansi, label_column= label_column)
                else:
                    logger.warning(f"Data Kosong untuk Nasional, Kategori: {kategori}")
        #Save PPT
        output_path = "output/ppt/Statistik_Nasional.pptx"
        prs.save(output_path)
        logger.info(f"PowerPoint Disimpan: {output_path}")
        #Export ke Excel
        '''if data_nasional and any(data_nasional.values()):
            export_per_provinsi(data_nasional)
        else:
            logger.info("Tidak ada data yang tersedia untuk diekspor.")'''
        #Convert JPG
        jpg_output_dir = f"output/jpg/"
        jpg_output_file = os.path.join(jpg_output_dir, "Statistik_Nasional.jpg")
        cmd = [
            "soffice",
            "--headless",
            "--convert-to","jpg",
            "--outdir", jpg_output_dir,
            output_path
        ]
        subprocess.run(cmd, check=True)
        logger.info(f"Slide diexport as JPG di {jpg_output_dir}")
        '''if os.path.isfile(output_path):
            os.remove(output_path)
            logger.info(f"File PPT {output_path} dihapus")'''
        sys.exit()

    #MODE INSTANSI
    #Mengambil Prefix Wilker
    filter_value = resolve_filter(args, conn)
    #Kondisi Mode
    if args.get("instansi"):
        mode = 'instansi'
        target_list = get_instansi_list(conn, filter_value)
    elif args.get("provinsi"):
        mode = 'provinsi'
        target_list = get_provinsi_prefix(conn, filter_value)
    elif args.get("kanreg"):
        mode = 'kanreg'
        target_list = get_kanreg_prefix(conn,filter_value)
    elif args.get("all"):
        mode = 'all'
        target_list = get_all_instansi(conn)
    elif args.get("allprovinsi"):
        mode = 'allprovinsi'
        target_list = get_all_provinsi(conn)
    else:
        raise ValueError("Mode harus 'instansi', 'provinsi', atau 'kanreg'")
    nama_table = get_table_name()
    logger.info(f"Nama tabel yang digunakan: {nama_table}")
    data_dict = get_all_data(conn, nama_table, target_list or [], mode=mode)
    # results = []
    for label, kategori_data in data_dict.items():
        prs = Presentation('template/template_infografis.pptx')
        logger.info(f"Data siap untuk Instansi: {label}")
        if mode in ("all","allprovinsi"):
            kategori_items = list(data_dict[label]["kategori"].items())
        else:
            kategori_items = list(kategori_data.items())
        for kategori, df in kategori_items:
            if 'count' in df.columns:
                #Penambahan Kolom Total pada Jenis ASN
                if 'jenis_asn' in df.columns and 'TOTAL' not in df['jenis_asn'].values:
                    df_total = pd.DataFrame({
                        'jenis_asn': ['TOTAL'],
                        'count': [df['count'].sum()]
                    })
                    df = pd.concat([df, df_total], ignore_index=True)
                    df_clean = df[df['jenis_asn'] != 'TOTAL'].copy()
                    df_clean = tambah_persentase(df_clean) #tambah persentase
                    df_final = pd.concat([
                        df_clean,
                        df[df['jenis_asn'] == 'TOTAL']
                    ], ignore_index=True)
                    if mode == "all":
                        data_dict[label]["kategori"][kategori] = df_final
                    else:
                        data_dict[label][kategori] = df_final
                #Penambahan Ringkasan Kelompok Jabatan       
                elif 'kelompok_jabatan' in df.columns and 'count' in df.columns:
                    df_ringkasan = ringkasan_kelompok_jabatan(df)
                    if mode in ("all","allprovinsi") :
                        data_dict[label]["kategori"][f"Ringkasan_{kategori}"] = df_ringkasan
                    else : 
                        data_dict[label][f"Ringkasan_{kategori}"] = df_ringkasan
                else:
                    if mode in ("all","allprovinsi") :
                        data_dict[label]["kategori"][kategori] = tambah_persentase(df)
                    else :
                        data_dict[label][kategori] = tambah_persentase(df)
        #Add Ulang Kategori  
        if mode in ("all","allprovinsi"):         
            kategori_items = list(data_dict[label]["kategori"].items())
        else:
            kategori_items = list(kategori_data.items())
        for kategori, df in kategori_items:
            if kategori == "Pendidikan":
                slide_index = selid_kategori_shared.get(kategori)
                slide = prs.slides[slide_index]
                if mode in ("all","allprovinsi"):
                    df_export = data_dict[label]["kategori"][kategori]
                else: 
                    df_export = data_dict[label][kategori]
                isi_balok_pendidikan(df_export, slide, mode)
            #Isi Data PPT
            if kategori in mapping_placeholder_per_kategori:
                mapping_placeholder = mapping_placeholder_per_kategori[kategori]
                label_column = label_column_per_kategori.get(kategori, "label")
                #Nama Instansi Dinamis                  
                mapping_instansi = generate_instansi_values(label)
                if mode in ("all","allprovinsi"):
                    data_export = data_dict[label]["kategori"][kategori]
                else:
                    data_export = data_dict[label][kategori]
                if data_export is not None:
                    isi_infografis(prs, data_export, mapping_placeholder, mapping_instansi, label_column= label_column)
                else:
                    logger.warning(f"Data Kosong untuk Provinsi: {label}, Kategori: {kategori}")
        #Save PPT
        safe_label = sanitize_filename(label.strip())
        output_path = f"output/ppt/{safe_label}.pptx"
        prs.save(output_path)
        logger.info(f"PowerPoint Disimpan: {output_path}")

        #Convert JPG
        if mode in ("all","allprovinsi"):
            instansi_id = data_dict[label]["instansi_id"]
        periode = datetime.now().strftime("%Y%m")
        periode_db = datetime.now().replace(day=1).date()
        if mode in ("all","allprovinsi"):
            jpg_output_dir = f"output/jpg/{instansi_id}"  
        else:
            jpg_output_dir = f"output/jpg/"
        cmd = [
            "soffice",
            "--headless",
            "--convert-to","jpg",
            "--outdir", jpg_output_dir,
            output_path
        ]
        subprocess.run(cmd, check=True)
        #Rename JPG
        if mode in ("all","allprovinsi"):
            default_jpg = os.path.join(jpg_output_dir, f"{safe_label}.jpg")
            rename_jpg = os.path.join(jpg_output_dir, f"{periode}.jpg")
            if os.path.exists(default_jpg):
                os.rename(default_jpg, rename_jpg)
                '''logger.info(f"Slide diexport as JPG di {rename_jpg}")'''
            #S3
            local_file = rename_jpg
            if mode == "all":
                s3_path = f"/instansi/{instansi_id}/{periode}.jpg"
            elif mode == "allprovinsi":
                s3_path = f"/provinsi/{instansi_id}/{periode}.jpg"
            s3_url = upload_file_s3(local_file, s3_path)
            logger.info(f"File tersimpan di: {s3_url}")
            save_to_db(periode_db, instansi_id, s3_path)
        else:
            logger.info(f"Slide diexport as JPG di {jpg_output_dir}")
        #Hapus PPT
        if mode in ("all","allprovinsi"):
            if os.path.isfile(output_path):
                os.remove(output_path)
                logger.info(f"File PPT {output_path} dihapus")

        #return hasil
        '''results.append({
        "status": "success",
        "mode": mode,
        "output_file": output_path
        })'''

         #Upload Drive
        '''SCOPES = ['https://www.googleapis.com/auth/drive.file']
        TOKEN_FILE = 'token.json'
        CREDENTIALS_FILE = 'credentials.json'
        if os.path.exists(TOKEN_FILE):
            creds = Credentials.from_authorized_user_file(TOKEN_FILE, SCOPES)
        else:
            flow = InstalledAppFlow.from_client_secrets_file(CREDENTIALS_FILE, SCOPES)
            creds = flow.run_local_server(port=0)
            with open(TOKEN_FILE, 'w') as token:
                token.write(creds.to_json())
        drive_service = build('drive', 'v3', credentials=creds)
        #save to bytes sementara
        ppt_stream = BytesIO()
        prs.save(ppt_stream)
        ppt_stream.seek(0)
        #Upload
        file_metadata = {
            'name': f'{label}.pptx',
            'parents': ['1yYlj9_Eq5YzRfinbqKUOYzVuWHfoebWq'] #1OhAKH94lYQwZVAarrz1nlWW86X3aBKJv
        }
        media = MediaIoBaseUpload(ppt_stream, mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation',
                              resumable=True)
        file = drive_service.files().create(body=file_metadata, media_body=media, fields='id').execute()
        print(f"File {label}.pptx diunggah ke Google Drive dengan ID: {file.get('id')}")'''

        #Export ke Excel
        '''if data_dict and any(data_dict.values()):
            export_per_provinsi(data_dict)
        else:
            print("Tidak ada data yang tersedia untuk diekspor.")'''
        
    '''return results'''


        
       